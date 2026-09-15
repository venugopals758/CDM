import os
from datetime import datetime

import openpyxl
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.db import connections
from django.shortcuts import render
from usermanagement.decorators import group_required
from usermanagement.models_gitam import EmployeeMaster, NirfData

PATENTS_MASTER_DATA_PATH = os.path.join(settings.BASE_DIR, 'static', 'patents_master_data.xlsx')


def _parse_patent_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    try:
        return datetime.strptime(value, '%d-%m-%Y').date()
    except (ValueError, TypeError):
        return None


def load_patents_master_data():
    wb = openpyxl.load_workbook(PATENTS_MASTER_DATA_PATH, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    col = {name: i for i, name in enumerate(header)}

    records = []
    for r in rows:
        empid = r[col['EmpID']]
        if empid is None:
            continue
        records.append({
            'empid': str(empid),
            'published_date': _parse_patent_date(r[col['Application Published Date']]),
            'grant_date': _parse_patent_date(r[col['Grant Date']]),
        })
    return records


def get_patent_counts(patents, emp_ids, year):
    emp_id_set = set(emp_ids)
    published = 0
    granted = 0
    for rec in patents:
        if rec['empid'] not in emp_id_set:
            continue
        if rec['published_date'] and rec['published_date'].year == year:
            published += 1
        if rec['grant_date'] and rec['grant_date'].year == year:
            granted += 1
    return published, granted

EMPLOYEE_CATEGORY_SQL = """
    SELECT EMPID, DOJ, DOR, 'Final' AS Category
    FROM EMPLOYEE_MASTER
    WHERE YEAR(DOJ) <= %s
      AND (DOR IS NULL OR YEAR(DOR) <= 1900 OR YEAR(DOR) > %s)
      AND JOB_TYPE = 'RE'
      AND EMP_TYPE LIKE 'T_'
    ORDER BY Category, EMPID
"""


SQL_PARAM_BATCH_SIZE = 1000


def _chunked(items, size):
    for i in range(0, len(items), size):
        yield items[i:i + size]


def get_phd_users(cursor, emp_ids):
    if not emp_ids:
        return []

    empid_list = []
    for batch in _chunked(emp_ids, SQL_PARAM_BATCH_SIZE):
        placeholders = ', '.join(['%s'] * len(batch))

        cursor.execute(
            f"""
            SELECT EMPID
            FROM td_EmpEduDetails
            WHERE QualificationName LIKE %s
              AND EMPID IN ({placeholders})
            """,
            ['%Ph.D.%'] + batch,
        )
        empid_list.extend(row[0] for row in cursor.fetchall())

        cursor.execute(
            f"""
            SELECT EMPID
            FROM EMPLOYEE_MASTER
            WHERE Highest_Qualification LIKE %s
              AND EMPID IN ({placeholders})
            """,
            ['%Ph.%'] + batch,
        )
        empid_list.extend(row[0] for row in cursor.fetchall())

    return empid_list


def get_experience_years_map(cursor, emp_ids, as_of_year):
    if not emp_ids:
        return {}

    earliest_start = {}
    for batch in _chunked(emp_ids, SQL_PARAM_BATCH_SIZE):
        placeholders = ', '.join(['%s'] * len(batch))

        cursor.execute(
            f"""
            SELECT EMPID, MIN(START_DATE) AS EARLIEST_START
            FROM (
                SELECT EMPID, [first] AS START_DATE
                FROM profiles_empworkexperience
                WHERE STATUS <> -1 AND EMPID IN ({placeholders})
                UNION ALL
                SELECT EMPID, FROM_DATE AS START_DATE
                FROM profiles_empingitamexperience
                WHERE STATUS <> -1 AND EMPID IN ({placeholders})
            ) AS combined_experience
            GROUP BY EMPID
            """,
            batch + batch,
        )
        for empid, start_date in cursor.fetchall():
            if start_date is not None and (empid not in earliest_start or start_date < earliest_start[empid]):
                earliest_start[empid] = start_date

    return {
        empid: as_of_year - start_date.year
        for empid, start_date in earliest_start.items()
    }


@login_required
def nirf_data(request):
    years = [2022, 2023, 2024, 2025]

    patents = load_patents_master_data()

    employee_ds = []
    nirf_rows = []
    with connections['GITAM'].cursor() as cursor:
        for year in years:

            cursor.execute(EMPLOYEE_CATEGORY_SQL, [year, year])
            empids = list(dict.fromkeys(row[0] for row in cursor.fetchall()))

            employee_ds.append({
                'year': year,
                'empids': empids,
            })

            phd_empids = set(get_phd_users(cursor, empids))

            with connections['ACADEMIC_PROFILES'].cursor() as exp_cursor:
                experience_years_map = get_experience_years_map(exp_cursor, empids, year)

            combo_empids_map = {}
            combo_gender_counts = {}
            for batch in _chunked(empids, SQL_PARAM_BATCH_SIZE):
                for campus_val, school_val, dept_val, empid_val, gender_val in EmployeeMaster.objects.using('GITAM').filter(
                    empid__in=batch
                ).values_list('campus', 'college_code', 'dept_code', 'empid', 'gender'):
                    key = (campus_val, school_val, dept_val)
                    combo_empids_map.setdefault(key, []).append(empid_val)
                    counts = combo_gender_counts.setdefault(key, {'M': 0, 'F': 0})
                    if gender_val in counts:
                        counts[gender_val] += 1

            for (campus, school, department), combo_empids in combo_empids_map.items():
                gender_counts = combo_gender_counts[(campus, school, department)]
                with_phd = sum(1 for e in combo_empids if e in phd_empids)

                exp_lt_8 = exp_8_15 = exp_gt_15 = 0
                for e in combo_empids:
                    exp_years = experience_years_map.get(e)
                    if exp_years is None:
                        continue
                    if exp_years < 8:
                        exp_lt_8 += 1
                    elif exp_years <= 15:
                        exp_8_15 += 1
                    else:
                        exp_gt_15 += 1

                patent_published, patent_granted = get_patent_counts(patents, combo_empids, year)

                row = {
                    'year': year,
                    'school': school,
                    'campus': campus,
                    'department': department,
                    'empid_list': combo_empids,
                    'faculty_count_regular': len(combo_empids),
                    'male_faculty': gender_counts['M'],
                    'female_faculty': gender_counts['F'],
                    'faculty_with_phd': with_phd,
                    'faculty_without_phd': len(combo_empids) - with_phd,
                    'exp_lt_8': exp_lt_8,
                    'exp_8_15': exp_8_15,
                    'exp_gt_15': exp_gt_15,
                    'patent_published': patent_published,
                    'patent_granted': patent_granted,
                    # 'publication_count': '',
                    # 'ret_count': '',
                    # 'research_funding': '',
                    # 'consulting_grants': '',
                    # 'fdp': '',
                    # 'mdp': '',
                }
                nirf_rows.append(row)


                NirfData.objects.using('GITAM').update_or_create(
                    year=year, campus=campus, school=school, department=department,
                    defaults={
                        'empid_list': ','.join(combo_empids),
                        'faculty_count_regular': row['faculty_count_regular'],
                        'male_faculty': row['male_faculty'],
                        'female_faculty': row['female_faculty'],
                        'faculty_with_phd': row['faculty_with_phd'],
                        'faculty_without_phd': row['faculty_without_phd'],
                        'exp_lt_8': row['exp_lt_8'],
                        'exp_8_15': row['exp_8_15'],
                        'exp_gt_15': row['exp_gt_15'],
                        'patent_published': row['patent_published'] or None,
                        'patent_granted': row['patent_granted'] or None,
                        # 'publication_count': row['publication_count'] or None,
                        # 'ret_count': row['ret_count'] or None,
                        # 'research_funding': row['research_funding'] or None,
                        # 'consulting_grants': row['consulting_grants'] or None,
                        # 'fdp': row['fdp'] or None,
                        # 'mdp': row['mdp'] or None,
                    },
                )

    context = {
        'employee_ds': employee_ds,
        'nirf_rows': nirf_rows,
    }
    return render(request, 'program/nirf.html', context)
